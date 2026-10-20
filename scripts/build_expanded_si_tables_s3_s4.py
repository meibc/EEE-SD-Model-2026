from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCREEN = ROOT / "output/amis_joint_bootstrap_direct_v2_n500/state_screen_and_response_categories/primary_states.csv"
OUT = ROOT / "output/supplementary_tables/expanded_SI_tables_S3_S4.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=65, start=70, bottom=65, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def caption(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Supplementary Table {number} | {title}")
    r.bold = True
    r.font.size = Pt(10.5)


def add_table(doc, headers, rows, widths, font_size=8.0, center_cols=()):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.style = "Table Grid"
    repeat_header(tbl.rows[0])
    for cell, label, w in zip(tbl.rows[0].cells, headers, widths):
        cell.text = label
        width(cell, w)
        margins(cell)
        shade(cell, "17365D")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(font_size)
    for ri, values in enumerate(rows):
        row = tbl.add_row()
        prevent_row_split(row)
        for ci, (cell, value, w) in enumerate(zip(row.cells, values, widths)):
            cell.text = str(value)
            width(cell, w)
            margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ri % 2:
                shade(cell, "F2F6FA")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if ci in center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(font_size)
    return tbl


def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Note: ")
    r.bold = True
    p.add_run(text)
    for run in p.runs:
        run.font.size = Pt(7.2)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.60)
sec.right_margin = Inches(0.60)
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(9)

# Table S3: one row for every state retained in the primary comparison.
screen = pd.read_csv(SCREEN).sort_values("state")
s3_rows = []
for _, x in screen.iterrows():
    s3_rows.append([
        x["state"],
        f'{int(x["n_complete_cells"])}/{int(x["n_expected_cells"])}',
        f'{x["min_cell_neff"]:.1f}',
        f'{x["q25_cell_neff"]:.1f}',
        f'{x["median_cell_neff"]:.1f}',
        "Retained",
        "Retained" if bool(x["strict_precision_screen"]) else "Excluded",
    ])

caption(doc, "S3", "State inclusion and survey-precision summary")
add_table(
    doc,
    ["State", "Complete cells", "Minimum n_eff", "25th percentile n_eff", "Median n_eff", "Primary set", "Stricter set"],
    s3_rows,
    [0.58, 0.84, 0.91, 1.10, 0.91, 0.86, 0.86],
    font_size=7.8,
    center_cols=(0, 1, 2, 3, 4, 5, 6),
)
note(
    doc,
    "The primary screen required all 32 state-period-measure cells, a median cell-specific Kish effective sample size (n_eff) of at least 60, and a 25th-percentile n_eff of at least 40. The stricter sensitivity screen required all 32 cells, a 25th-percentile n_eff of at least 50, and a minimum n_eff of at least 25. The stricter screen retained 15 states; Massachusetts and Tennessee were excluded. Screening was specified independently of model fit and scenario response. State-level statistics for states not retained in the primary comparison are supplied in Supplementary Data 1.",
)

# Landscape section for detailed scenario definitions.
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
sec2.orientation = WD_ORIENT.LANDSCAPE
sec2.page_width, sec2.page_height = sec2.page_height, sec2.page_width
sec2.top_margin = Inches(0.48)
sec2.bottom_margin = Inches(0.48)
sec2.left_margin = Inches(0.48)
sec2.right_margin = Inches(0.48)

caption(doc, "S4", "Definitions and implementation of stigma-related scenarios")
s4_rows = [
    ["Baseline", "Reference", "None", "None", "None", "Observed 2024 state propagated through the fitted system without a scenario perturbation; comparator for paired contrasts."],
    ["Reduce anticipated healthcare stigma", "Primary proportional", "Anticipated healthcare stigma", "50% reduction from the draw-specific counterfactual baseline probability", "None", "Isolates the response to lowering healthcare-facing stigma."],
    ["Reduce general social stigma", "Additional proportional", "General social stigma", "50% reduction from the draw-specific counterfactual baseline probability", "None", "Examines a broader social-stigma level perturbation."],
    ["Reduce family stigma", "Additional proportional", "Family stigma", "50% reduction from the draw-specific counterfactual baseline probability", "None", "Examines a family-stigma level perturbation."],
    ["Weaken AHS → PrEP pathway", "Primary pathway", "None", "None", "Multiply the fitted AHS → PrEP coefficient by 0.50.", "Tests weakening a specific stigma–prevention linkage without changing stigma prevalence."],
    ["Weaken AHS → HIV-testing pathway", "Additional pathway", "None", "None", "Multiply the fitted AHS → HIV-testing coefficient by 0.50.", "Tests weakening the fitted connection between AHS and HIV testing."],
    ["Strengthen healthcare-contact → lower-AHS pathway", "Additional pathway", "None", "None", "Multiply the fitted healthcare-contact → AHS coefficient by 1.50.", "Intended to examine a stronger protective feedback from care contact to subsequent AHS."],
    ["Attenuate healthcare-generated AHS", "Additional pathway", "None", "None", "Reduce a positive healthcare-contact → AHS coefficient by 50%; nonpositive coefficients are unchanged.", "Examines mitigation where healthcare contact is associated with greater subsequent AHS."],
    ["Reverse healthcare-generated AHS", "Additional pathway", "None", "None", "Move a positive healthcare-contact → AHS coefficient to −0.087, the prespecified typical protective-state target.", "Exploratory reversal of a harmful healthcare-contact feedback."],
    ["Combined stigma package", "Primary proportional", "AHS, general social stigma, and family stigma", "Each stigma probability reduced by 50% from its draw-specific counterfactual baseline", "Multiply the fitted AHS → PrEP coefficient by 0.50.", "Assesses coordinated changes in stigma levels and a proximal stigma–PrEP pathway."],
    ["Fixed 10-point AHS reduction", "Primary fixed-dose comparison", "Anticipated healthcare stigma", "Subtract 0.10 from the draw-specific counterfactual baseline probability; lower bound 10⁻⁶", "None", "Uses a common absolute dose across states rather than a baseline-dependent proportional dose."],
    ["Fixed 10-point combined package", "Primary fixed-dose comparison", "AHS, general social stigma, and family stigma", "Subtract 0.10 from each draw-specific counterfactual baseline probability; lower bound 10⁻⁶", "Multiply the fitted AHS → PrEP coefficient by 0.50.", "Fixed-dose counterpart to the proportional combined package."],
    ["Fixed 5-point AHS reduction", "Dose sensitivity", "Anticipated healthcare stigma", "Subtract 0.05 from the draw-specific counterfactual baseline probability; lower bound 10⁻⁶", "None", "Evaluates sensitivity to a smaller common absolute dose."],
    ["Fixed 5-point combined package", "Dose sensitivity", "AHS, general social stigma, and family stigma", "Subtract 0.05 from each draw-specific counterfactual baseline probability; lower bound 10⁻⁶", "Multiply the fitted AHS → PrEP coefficient by 0.50.", "Smaller fixed-dose counterpart to the combined package."],
]
add_table(
    doc,
    ["Scenario", "Role in analysis", "State variable(s) modified", "Probability-scale change", "Pathway change", "Interpretive comparison"],
    s4_rows,
    [1.45, 1.12, 1.55, 2.25, 2.05, 2.10],
    font_size=6.7,
)
note(
    doc,
    "Stigma-level perturbations were applied on the probability scale before conversion to centered-logit coordinates. Targets were 50% of the draw-specific counterfactual baseline or that baseline minus 0.10/0.05. Effects ramped from 2026 through 2032 over three two-year intervals and were maintained through 2036. Paired dynamic-model and epidemiological draws were compared with their paired 2036 baselines. Scenarios are standardized perturbations of fitted systems, not causal effects of specific interventions. AHS, anticipated healthcare stigma; PrEP, pre-exposure prophylaxis.",
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
