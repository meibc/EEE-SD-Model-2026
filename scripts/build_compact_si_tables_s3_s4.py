from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "supplementary_tables" / "compact_SI_tables_S3_S4.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_caption(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Supplementary Table {number} | ")
    r.bold = True
    r2 = p.add_run(title)
    r2.bold = True
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, (cell, text, width) in enumerate(zip(header.cells, headers, widths)):
        cell.text = text
        set_cell_width(cell, width)
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
    for ri, values in enumerate(rows):
        row = table.add_row()
        for cell, value, width in zip(row.cells, values, widths):
            cell.text = str(value)
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ri % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(8.5)
    return table


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Note: ")
    r.bold = True
    p.add_run(text)
    for run in p.runs:
        run.font.size = Pt(8.5)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(16)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Compact Supplementary Methods Tables")
intro = doc.add_paragraph(
    "These compact tables are intended for inclusion in the Supplementary Information document. "
    "The complete state screen and model specification remain available in Supplementary Data 1."
)
intro.paragraph_format.space_after = Pt(8)

add_caption(doc, "S3", "State inclusion summary")
s3_headers = ["Analysis set", "Completeness requirement", "Precision requirement", "Retained states"]
s3_rows = [
    [
        "Primary comparison\n(n = 17)",
        "All 32 period-by-measure cells available",
        "Median cell-specific n_eff ≥60 and 25th percentile n_eff ≥40",
        "CA, FL, GA, IL, IN, MA, MD, MI, NC, NJ, NY, OH, PA, TN, TX, VA, and WA",
    ],
    [
        "Stricter sensitivity set\n(n = 15)",
        "All 32 period-by-measure cells available",
        "25th percentile n_eff ≥50 and minimum n_eff ≥25",
        "CA, FL, GA, IL, IN, MD, MI, NC, NJ, NY, OH, PA, TX, VA, and WA",
    ],
]
add_table(doc, s3_headers, s3_rows, [1.25, 1.75, 2.20, 3.25])
add_note(
    doc,
    "Each state contributed four two-year periods and eight modeled AMIS measures. Kish effective sample size was calculated separately for each state-period-measure cell. State inclusion depended only on completeness and survey precision, not on model fit or scenario response. Full state-level screening results are provided in Supplementary Data 1.",
)

doc.add_paragraph().paragraph_format.space_after = Pt(1)
add_caption(doc, "S4", "Definitions of the primary stigma-related scenarios")
s4_headers = ["Scenario", "Stigma-level perturbation", "Pathway perturbation", "Purpose"]
s4_rows = [
    [
        "Reduce anticipated healthcare stigma",
        "AHS was moved toward 50% of its draw-specific counterfactual baseline probability.",
        "None",
        "Isolated response to a proportional reduction in healthcare-facing stigma.",
    ],
    [
        "Weaken AHS → PrEP pathway",
        "None",
        "The fitted AHS → PrEP coefficient was attenuated by 50%.",
        "Response to weakening a specific stigma–prevention relationship without directly changing stigma prevalence.",
    ],
    [
        "Combined stigma package",
        "AHS, general social stigma, and family stigma were each moved toward 50% of their draw-specific baseline probabilities.",
        "The fitted AHS → PrEP coefficient was attenuated by 50%.",
        "Response to simultaneous changes across healthcare-facing, broader social, and family stigma processes.",
    ],
    [
        "Fixed 10-percentage-point scenarios",
        "AHS alone, or AHS, general social stigma, and family stigma together, were each reduced by 10 percentage points, bounded above zero.",
        "The combined fixed-dose scenario also attenuated the AHS → PrEP coefficient by 50%.",
        "Sensitivity analysis using a common absolute dose rather than a baseline-dependent proportional reduction.",
    ],
]
add_table(doc, s4_headers, s4_rows, [1.55, 2.55, 2.05, 2.30])
add_note(
    doc,
    "Perturbations began at the first two-year model step after the final observed period, increased linearly over three two-year steps, and were then maintained through 2036. Stigma-level changes were implemented on the probability scale before transformation to centered-logit coordinates. Baseline and scenario trajectories were evaluated using paired dynamic-model and epidemiological draws. These scenarios were standardized perturbations of fitted systems and were not interpreted as causal effects of specific real-world interventions. AHS, anticipated healthcare stigma; PrEP, pre-exposure prophylaxis.",
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
